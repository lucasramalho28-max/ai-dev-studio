import csv
import hashlib
import io
import json
import os
import re
import uuid
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

os.environ.setdefault('DISABLE_SQLALCHEMY_CEXT_RUNTIME', '1')

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import NullPool

from models.logistics import Base, LogisticsActivity, LogisticsAnalysis, LogisticsDashboard, LogisticsIndicator, LogisticsMemory, LogisticsUpload, utc_now


MAX_UPLOAD = 25 * 1024 * 1024
MAX_ROWS = 100_000
ALLOWED = {'.xlsx', '.csv', '.txt', '.json'}

ROLE_ALIASES = {
    'order_id': ('pedido','id_pedido','numero_pedido','n_pedido','order_id','ordem'),
    'branch': ('filial','cd','centro_distribuicao','centro_de_distribuicao','unidade','branch'),
    'regional': ('regional','regiao','region'),
    'week': ('semana','week','semana_ano','week_year'),
    'on_time': ('no_prazo','dentro_prazo','on_time','sla_ok','cumpriu_sla'),
    'status': ('status','situacao','status_pedido'),
    'delay_reason': ('motivo_atraso','causa_atraso','reason','motivo'),
    'due_date': ('data_limite','prazo','data_prometida','due_date','sla_data'),
    'ship_date': ('data_expedicao','data_entrega','data_realizada','ship_date','delivery_date'),
    'quantity': ('quantidade','qtd','volume','quantity'),
    'value': ('valor','valor_total','custo','value'),
    'address': ('endereco','endereco_armazenagem','location'),
    'product': ('produto','sku','codigo_produto','product'),
}

CATALOG = [
    dict(key='sla_expedicao',name='SLA de expedição',description='Percentual de pedidos expedidos dentro do prazo conforme regra aprovada.',formula='pedidos únicos no prazo / pedidos únicos elegíveis × 100',fields=['order_id','branch','on_time ou due_date + ship_date'],unit='%',filters=['período','filial','regional']),
    dict(key='acuracidade',name='Acuracidade de inventário',description='Índice de aderência entre estoque físico e sistêmico.',formula='Configurável: metodologia deve ser informada pelo usuário',fields=['estoque_fisico','estoque_sistema','chave do item'],unit='%',filters=['filial','endereço','ciclo']),
    dict(key='avarias',name='Avarias',description='Quantidade e valor financeiro de itens avariados.',formula='soma da quantidade e/ou valor dos registros classificados como avaria',fields=['quantity','value','branch','motivo'],unit='unidades / R$',filters=['período','filial','motivo']),
    dict(key='aging',name='Aging',description='Distribuição de pendências por faixas de dias.',formula='data de referência − data de início da pendência',fields=['data_inicio','status','chave'],unit='dias',filters=['faixa','filial','status']),
    dict(key='backlog',name='Backlog',description='Pedidos pendentes conforme status aprovados.',formula='contagem distinta de pedidos nos status definidos como pendentes',fields=['order_id','status'],unit='pedidos',filters=['filial','status','período']),
    dict(key='produtividade',name='Produtividade',description='Volume produzido em relação ao recurso ou período disponível.',formula='volume aprovado / recurso ou período aprovado',fields=['quantity','recurso ou período'],unit='configurável',filters=['filial','turno','período']),
]


def normalize(value):
    import unicodedata
    text = unicodedata.normalize('NFKD', str(value)).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','_',text).strip('_')


class DataTable:
    def __init__(self, rows, columns=None):
        self.rows = rows
        self.columns = columns or list(rows[0]) if rows else (columns or [])

    def __len__(self):
        return len(self.rows)


def missing(value):
    return value is None or (isinstance(value, str) and not value.strip())


def number(value):
    if missing(value) or isinstance(value, bool):
        return None
    try:
        text = str(value).strip().replace(' ', '')
        if ',' in text and '.' in text:
            text = text.replace('.', '').replace(',', '.')
        elif ',' in text:
            text = text.replace(',', '.')
        return float(text)
    except (ValueError, TypeError):
        return None


def date_value(value):
    if isinstance(value, datetime):
        return value
    if missing(value):
        return None
    text = str(value).strip()
    for pattern in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y', '%d/%m/%Y %H:%M:%S', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            pass
    return None


class LogisticsService:
    def __init__(self, data_root, store):
        self.root = Path(data_root).resolve() / 'logistics'
        self.root.mkdir(parents=True, exist_ok=True)
        self.store = store
        self.engine = create_engine(f'sqlite:///{(Path(data_root)/"studio.sqlite3").as_posix()}',
                                    connect_args={'check_same_thread':False}, poolclass=NullPool)
        Base.metadata.create_all(self.engine)

    def _project_root(self, project):
        if not re.fullmatch(r'[a-f0-9]{32}', project):
            raise ValueError('Projeto inválido')
        path = (self.root / project).resolve()
        if not path.is_relative_to(self.root):
            raise ValueError('Diretório de dados inválido')
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _activity(self, project, kind, detail):
        with Session(self.engine) as db:
            db.add(LogisticsActivity(project=project,kind=kind,detail=detail))
            db.commit()

    def save_upload(self, project, uploaded):
        original = Path(uploaded.filename or '').name
        extension = Path(original).suffix.lower()
        if extension not in ALLOWED:
            raise ValueError('Formato permitido: .xlsx, .csv, .txt ou .json')
        raw = uploaded.stream.read(MAX_UPLOAD + 1)
        if not raw or len(raw) > MAX_UPLOAD:
            raise ValueError('Arquivo vazio ou maior que 25 MB')
        if extension == '.xlsx':
            self._validate_xlsx(raw)
        uid = uuid.uuid4().hex
        stored = uid + extension
        upload_dir = self._project_root(project) / 'uploads'
        upload_dir.mkdir(exist_ok=True)
        target = upload_dir / stored
        target.write_bytes(raw)
        digest = hashlib.sha256(raw).hexdigest()
        with Session(self.engine) as db:
            db.add(LogisticsUpload(id=uid,project=project,original_name=original[:180],stored_name=stored,
                                   extension=extension,size=len(raw),sha256=digest))
            db.commit()
        self._activity(project,'upload',f'Base local carregada: {original} ({len(raw)} bytes)')
        return {'id':uid,'name':original,'size':len(raw),'sha256':digest,'extension':extension}

    @staticmethod
    def _validate_xlsx(raw):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 3000 or sum(m.file_size for m in members) > 250 * 1024 * 1024:
                    raise ValueError('Planilha compactada excede os limites seguros')
                if any('vbaProject' in m.filename or m.filename.lower().endswith('.bin') for m in members):
                    raise ValueError('Macros e conteúdo binário não são aceitos')
        except zipfile.BadZipFile:
            raise ValueError('Arquivo XLSX inválido') from None

    def upload(self, project, upload_id):
        with Session(self.engine) as db:
            item = db.get(LogisticsUpload, upload_id)
            if not item or item.project != project:
                raise ValueError('Base não encontrada neste projeto')
            return {'id':item.id,'project':item.project,'name':item.original_name,'stored':item.stored_name,
                    'extension':item.extension,'size':item.size,'sha256':item.sha256,'created':item.created}

    def _read(self, project, upload_id):
        meta = self.upload(project, upload_id)
        path = self._project_root(project) / 'uploads' / meta['stored']
        ext = meta['extension']
        if ext == '.xlsx':
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, ())
            columns = [str(value).strip() if value is not None else '' for value in header]
            if not columns or any(not name for name in columns):
                raise ValueError('A primeira linha da planilha deve conter nomes de colunas')
            rows = []
            for values in iterator:
                if len(rows) > MAX_ROWS:
                    break
                rows.append({column: values[index] if index < len(values) else None for index, column in enumerate(columns)})
            workbook.close()
        elif ext == '.json':
            value = json.loads(path.read_text(encoding='utf-8-sig'))
            if isinstance(value, dict):
                for key in ('data', 'rows', 'records', 'items'):
                    if isinstance(value.get(key), list):
                        value = value[key]
                        break
            if isinstance(value, dict):
                value = [value]
            if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
                raise ValueError('JSON deve conter um objeto ou uma lista de objetos')
            columns = []
            for item in value:
                for key in item:
                    name = str(key).strip()
                    if name not in columns:
                        columns.append(name)
            rows = [{column: item.get(column) for column in columns} for item in value[:MAX_ROWS + 1]]
        else:
            text = self._decode(path.read_bytes())
            if ext == '.txt' and not any(x in text[:4096] for x in (',', ';', '\t', '|')):
                columns = ['texto']
                rows = [{'texto': line} for line in text.splitlines() if line.strip()][:MAX_ROWS + 1]
            else:
                try:
                    sep = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|').delimiter
                except csv.Error:
                    sep = ','
                reader = csv.DictReader(io.StringIO(text), delimiter=sep)
                columns = [str(column).strip() for column in (reader.fieldnames or [])]
                rows = []
                for item in reader:
                    if len(rows) > MAX_ROWS:
                        break
                    rows.append({column: item.get(column) for column in columns})
        if len(rows) > MAX_ROWS:
            raise ValueError(f'Limite de análise: {MAX_ROWS:,} linhas')
        columns = [str(column).strip()[:120] for column in columns]
        if not columns:
            raise ValueError('A base não possui colunas identificáveis')
        if len(set(columns)) != len(columns):
            raise ValueError('A base possui nomes de colunas duplicados')
        return DataTable(rows, columns), meta

    @staticmethod
    def _decode(raw):
        for encoding in ('utf-8-sig','utf-8','cp1252','latin1'):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                pass
        raise ValueError('Não foi possível identificar a codificação do texto')

    def analyze(self, project, upload_id):
        table, meta = self._read(project, upload_id)
        normalized = {normalize(column): column for column in table.columns}
        mapping = {}
        for role, aliases in ROLE_ALIASES.items():
            alias_set = {normalize(alias) for alias in aliases}
            hits = [original for key, original in normalized.items() if key in alias_set]
            if len(hits) == 1:
                mapping[role] = hits[0]
        columns = []
        for name in table.columns:
            values = [row.get(name) for row in table.rows]
            present = [value for value in values if not missing(value)]
            numeric_ratio = sum(number(value) is not None for value in present) / max(len(present), 1)
            date_ratio = sum(date_value(value) is not None for value in present) / max(len(present), 1)
            boolean_values = {normalize(value) for value in present}
            if numeric_ratio >= .85:
                kind = 'número'
            elif date_ratio >= .85:
                kind = 'data'
            elif boolean_values <= {'true', 'false', 'sim', 'nao', '1', '0', 'yes', 'no'}:
                kind = 'booleano'
            else:
                kind = 'texto'
            missing_count = len(values) - len(present)
            serialized = {json.dumps(value, ensure_ascii=False, sort_keys=True, default=str) for value in present}
            columns.append({'name': name, 'type': kind, 'missing': missing_count,
                            'missing_pct': round(100 * missing_count / max(len(values), 1), 2),
                            'unique': len(serialized)})
        signatures = [tuple(json.dumps(row.get(column), ensure_ascii=False, sort_keys=True, default=str) for column in table.columns) for row in table.rows]
        duplicate_rows = sum(count - 1 for count in Counter(signatures).values() if count > 1)
        key_duplicates = None
        if mapping.get('order_id'):
            keys = [str(row.get(mapping['order_id'])) for row in table.rows if not missing(row.get(mapping['order_id']))]
            key_duplicates = sum(count for count in Counter(keys).values() if count > 1)
        warnings = []
        if duplicate_rows:
            warnings.append(f'{duplicate_rows} linha(s) integralmente duplicada(s).')
        missing_columns = [column['name'] for column in columns if column['missing']]
        if missing_columns:
            warnings.append('Valores ausentes em: ' + ', '.join(missing_columns[:12]) + '.')
        if key_duplicates:
            warnings.append(f'{key_duplicates} registro(s) compartilham a mesma chave de pedido; confirme a granularidade.')
        sla_ready = bool(mapping.get('order_id') and mapping.get('branch') and
                         (mapping.get('on_time') or (mapping.get('due_date') and mapping.get('ship_date'))))
        questions = []
        if not mapping.get('order_id'):
            questions.append('Qual coluna identifica unicamente o pedido?')
        if not mapping.get('branch'):
            questions.append('Qual coluna representa a filial ou CD?')
        if not mapping.get('on_time') and not (mapping.get('due_date') and mapping.get('ship_date')):
            questions.append('Como o prazo deve ser validado: coluna pronta de SLA ou comparação entre quais datas?')
        profile = {'analysis_id': uuid.uuid4().hex, 'upload': meta, 'rows': len(table),
                   'columns_count': len(table.columns), 'columns': columns,
                   'duplicate_rows': duplicate_rows, 'duplicate_order_records': key_duplicates,
                   'mapping': mapping, 'sla_ready': sla_ready, 'warnings': warnings, 'questions': questions,
                   'suggested_indicators': [item['key'] for item in CATALOG if self._indicator_available(item, mapping)]}
        with Session(self.engine) as db:
            db.add(LogisticsAnalysis(id=profile['analysis_id'], project=project, upload_id=upload_id,
                                     profile_json=json.dumps(profile, ensure_ascii=False)))
            db.commit()
        summary = json.dumps({'base': meta['name'], 'rows': len(table), 'columns': table.columns,
                              'mapping': mapping, 'warnings': warnings, 'questions': questions}, ensure_ascii=False)
        self.store.set('logistics_context:' + project, summary)
        self._activity(project, 'analysis', f'Base analisada: {meta["name"]}; {len(table)} linhas; {len(table.columns)} colunas')
        return profile

    @staticmethod
    def _indicator_available(indicator, mapping):
        if indicator['key'] == 'sla_expedicao':
            return bool(mapping.get('order_id') and mapping.get('branch'))
        if indicator['key'] == 'avarias':
            return bool(mapping.get('quantity') or mapping.get('value'))
        if indicator['key'] == 'backlog':
            return bool(mapping.get('order_id') and mapping.get('status'))
        return False

    def latest_analysis(self, project, upload_id):
        with Session(self.engine) as db:
            item = db.scalars(select(LogisticsAnalysis).where(LogisticsAnalysis.project==project,LogisticsAnalysis.upload_id==upload_id).order_by(LogisticsAnalysis.created.desc())).first()
            return json.loads(item.profile_json) if item else self.analyze(project,upload_id)

    def state(self, project):
        with Session(self.engine) as db:
            uploads = db.scalars(select(LogisticsUpload).where(LogisticsUpload.project==project).order_by(LogisticsUpload.created.desc())).all()
            dashboards = db.scalars(select(LogisticsDashboard).where(LogisticsDashboard.project==project).order_by(LogisticsDashboard.created.desc())).all()
            activities = db.scalars(select(LogisticsActivity).where(LogisticsActivity.project==project).order_by(LogisticsActivity.id.desc()).limit(20)).all()
            configured = db.scalars(select(LogisticsIndicator).where(LogisticsIndicator.project==project)).all()
            memory = db.scalars(select(LogisticsMemory).where(LogisticsMemory.project==project,LogisticsMemory.approved==True)).all()
        return {
            'uploads':[{'id':x.id,'name':x.original_name,'size':x.size,'sha256':x.sha256,'created':x.created} for x in uploads],
            'dashboards':[{'id':x.id,'title':x.title,'status':x.status,'created':x.created,'result':json.loads(x.result_json)} for x in dashboards],
            'activities':[{'kind':x.kind,'detail':x.detail,'created':x.created} for x in activities],
            'catalog':CATALOG,
            'configured_indicators':[self._indicator_dict(x) for x in configured],
            'memory':[{'category':x.category,'key':x.key,'value':json.loads(x.value_json),'updated':x.updated} for x in memory],
        }

    @staticmethod
    def _indicator_dict(x):
        return {'id':x.id,'key':x.key,'name':x.name,'description':x.description,'formula':x.formula,
                'fields':json.loads(x.fields_json),'filters':json.loads(x.filters_json),'unit':x.unit,
                'target':x.target,'source':x.source,'updated':x.updated,'approved':x.approved}

    def save_indicator(self, project, payload):
        required = ('name','description','formula','unit','source')
        if any(not str(payload.get(k,'')).strip() for k in required):
            raise ValueError('Nome, descrição, fórmula, unidade e fonte são obrigatórios')
        uid = payload.get('id') or uuid.uuid4().hex
        with Session(self.engine) as db:
            item = db.get(LogisticsIndicator,uid)
            if item and item.project != project: raise ValueError('Indicador pertence a outro projeto')
            if not item:
                item=LogisticsIndicator(id=uid,project=project,key=normalize(payload.get('key') or payload['name']))
                db.add(item)
            item.name=str(payload['name'])[:120]; item.description=str(payload['description'])[:2000]
            item.formula=str(payload['formula'])[:2000]; item.fields_json=json.dumps(payload.get('fields',[]),ensure_ascii=False)
            item.filters_json=json.dumps(payload.get('filters',[]),ensure_ascii=False); item.unit=str(payload['unit'])[:40]
            item.target=str(payload.get('target',''))[:80]; item.source=str(payload['source'])[:180]
            item.approved=True; item.updated=utc_now(); db.commit()
        self._activity(project,'indicator',f'Indicador aprovado: {payload["name"]}')
        return {'id':uid,'approved':True}

    def save_memory(self, project, category, key, value):
        allowed={'branches','regions','formula','filter','target','presentation','schema'}
        if category not in allowed or not str(key).strip(): raise ValueError('Categoria ou chave de memória inválida')
        uid=uuid.uuid4().hex
        with Session(self.engine) as db:
            db.add(LogisticsMemory(id=uid,project=project,category=category,key=str(key)[:100],
                                   value_json=json.dumps(value,ensure_ascii=False),approved=True))
            db.commit()
        self._activity(project,'memory',f'Memória aprovada: {category}/{key}')
        return {'id':uid,'approved':True}

    def synthetic_sla(self, project):
        rows=[]
        branches=['CD Norte','CD Centro','CD Sul','CD Leste']
        reasons=['Transportadora','Separação','Estoque','Documentação']
        for week in range(1,9):
            for branch_index,branch in enumerate(branches):
                for order in range(1,26):
                    on_time=((order+week+branch_index*2)%10) >= branch_index+1
                    rows.append({'pedido':f'S{week:02d}{branch_index}{order:03d}','filial':branch,
                                 'regional':['Norte','Centro-Oeste','Sul','Sudeste'][branch_index],
                                 'semana':f'2026-S{week:02d}','no_prazo':'Sim' if on_time else 'Não',
                                 'motivo_atraso':'' if on_time else reasons[(order+week)%len(reasons)]})
        data=io.StringIO(newline='')
        writer=csv.DictWriter(data,fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
        class Upload:
            filename='dados_sinteticos_sla.csv'
            stream=io.BytesIO(data.getvalue().encode('utf-8'))
        saved=self.save_upload(project,Upload())
        self._activity(project,'synthetic','Base sintética criada; não representa dados reais da empresa')
        return saved
